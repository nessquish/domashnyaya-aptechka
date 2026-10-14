"""Экспорт и импорт аптечки, полный сброс данных пользователя."""

import csv
import json
import re
import shutil
import stat
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Dict, List, Optional, Tuple, Union

from pharmacy.db.connection import Database
from pharmacy.errors import ValidationError
from pharmacy.models import ShoppingSource
from pharmacy.services.product_service import ProductForm, ProductService
from pharmacy.services.shopping_service import ShoppingService
from pharmacy.utils.dates import format_user_date
from pharmacy.utils.formatting import format_quantity

FORMATS = ("csv", "json", "xlsx")
CSV_DELIMITER = ";"  # так открывают файл русские версии Excel

# Ключ -> заголовок столбца в CSV и Excel.
COLUMNS: Dict[str, str] = {
    "name": "Название",
    "category": "Категория",
    "quantity": "Количество",
    "unit": "Ед. изм.",
    "expiry_date": "Срок годности",
    "indications": "Показания",
    "storage_place": "Место хранения",
    "min_quantity": "Мин. остаток",
    "note": "Примечание",
}
TEMPLATE_XLSX = Path(__file__).resolve().parent.parent / "assets" / "template.xlsx"
KIT_SHEET = "Моя аптечка"  # листы шаблона
SHOPPING_SHEET = "Список покупок"
SHOPPING_TITLE = "Товар"  # на листе покупок это название позиции
HEADER_SEARCH_ROWS = 5  # над заголовками шаблона есть строка с названием таблицы


def _norm(title: object) -> str:
    """Заголовок для сравнения: без регистра, пробелов и точек."""
    return re.sub(r"[\s.]", "", str(title).casefold())


_BY_TITLE = {_norm(title): key for key, title in COLUMNS.items()}


def _table_records(rows: List[tuple]) -> List[Dict[str, object]]:
    """Строки таблицы листа как словари «ключ столбца -> значение».

    Заголовки ищутся в первых строках листа. Если их нет, таблицы нет.
    """
    keys_by_title = {**_BY_TITLE, _norm(SHOPPING_TITLE): "name"}

    def keys_of(row: tuple) -> List[Optional[str]]:
        return [
            keys_by_title.get(_norm(cell)) if cell is not None else None for cell in row
        ]

    header_at = next(
        (
            index
            for index, row in enumerate(rows[:HEADER_SEARCH_ROWS])
            if sum(key is not None for key in keys_of(row)) >= 2
        ),
        None,
    )
    if header_at is None:
        return []
    keys = keys_of(rows[header_at])
    records = []
    for number, raw in enumerate(rows[header_at + 1 :], start=header_at + 2):
        record = {
            key: value
            for key, value in zip(keys, raw)
            if key is not None and value is not None
        }
        if record:
            records.append({**record, "_row": number})
    return records


def _cell_text(value: object) -> str:
    """Значение ячейки как текст формы (дата из Excel становится ДД.ММ.ГГГГ)."""
    if isinstance(value, datetime):
        value = value.date()
    if isinstance(value, date):
        return format_user_date(value)
    return "" if value is None else str(value).strip()


@dataclass
class ImportResult:
    """Итог импорта.

    Attributes:
        added: Сколько товаров добавлено в аптечку.
        shopping_added: Сколько позиций добавлено в список покупок.
        issues: Пропущенные строки: (лист, номер строки, причина).
    """

    added: int = 0
    shopping_added: int = 0
    issues: List[Tuple[str, int, str]] = field(default_factory=list)

    @property
    def errors(self) -> List[str]:
        """Пропущенные строки текстом («Моя аптечка, строка 3: причина»)."""
        return [f"{sheet}, строка {row}: {why}" for sheet, row, why in self.issues]

    @property
    def skipped(self) -> int:
        """Сколько строк пропущено."""
        return len(self.issues)

    def report(self) -> str:
        """Итог для показа: сколько добавлено и все причины пропуска.

        Одинаковые причины собраны вместе: «Ошибка импорта: причина (лист,
        строки 3, 5)».
        """
        lines = [
            f"Добавлено товаров в аптечку: {self.added}",
            f"Добавлено товаров в список покупок: {self.shopping_added}",
        ]
        reasons: Dict[str, Dict[str, List[int]]] = {}
        for sheet, row, why in self.issues:
            reasons.setdefault(why, {}).setdefault(sheet, []).append(row)
        for why, sheets in reasons.items():
            places = "; ".join(
                f"{sheet}, {'строка' if len(rows) == 1 else 'строки'} "
                + ", ".join(map(str, rows))
                for sheet, rows in sheets.items()
            )
            lines.append(f"Ошибка импорта: {why} ({places})")
        return "\n".join(lines)


def free_target(path: Union[str, Path]) -> None:
    """Готовит существующий файл к замене.

    Снимает отметку «только чтение». Если файл занят другой программой
    (например, открыт в Excel), понятно говорит об этом.

    Raises:
        ValidationError: Если файл нельзя перезаписать.
    """
    path = Path(path)
    if not path.is_file():
        return
    try:
        path.chmod(stat.S_IREAD | stat.S_IWRITE)
    except OSError:
        pass
    try:
        path.open("ab").close()
    except OSError:
        raise ValidationError(
            "Файл нельзя заменить: он открыт в другой программе (например, "
            "в Excel) или защищён от записи. Закройте его и сохраните снова."
        )


class DataService:
    """Работа с данными пользователя целиком: файлы и сброс."""

    def __init__(
        self,
        db: Database,
        products: ProductService,
        shopping: Optional[ShoppingService] = None,
    ) -> None:
        """Создаёт сервис.

        Args:
            db: Менеджер базы данных.
            products: Сервис товаров (импорт идёт через его проверки).
            shopping: Сервис списка покупок (для правой таблицы шаблона).
        """
        self._db = db
        self._products = products
        self._shopping = shopping

    # --- экспорт ---

    def export_products(self, user_id: int, path: Union[str, Path], fmt: str) -> int:
        """Записывает все товары пользователя в файл.

        Args:
            user_id: Владелец товаров.
            path: Куда сохранить.
            fmt: ``csv``, ``json`` или ``xlsx``.

        Returns:
            Сколько товаров записано.

        Raises:
            ValidationError: Если формат неизвестен.
        """
        if fmt not in FORMATS:
            raise ValidationError("Выберите формат файла: CSV, JSON или Excel")
        rows = self._rows(user_id)
        path = Path(path)
        free_target(path)
        if fmt == "json":
            path.write_text(
                json.dumps(rows, ensure_ascii=False, indent=2), encoding="utf-8"
            )
        elif fmt == "csv":
            with path.open("w", encoding="utf-8-sig", newline="") as handle:
                writer = csv.writer(handle, delimiter=CSV_DELIMITER)
                writer.writerow(COLUMNS.values())
                for row in rows:
                    writer.writerow(self._cells(row))
        else:
            self._write_excel(path, rows)
        return len(rows)

    def _rows(self, user_id: int) -> List[Dict[str, object]]:
        rows = []
        for view in self._products.list_products(user_id, sort="name"):
            product = view.product
            special = {
                "category": product.category_name,
                "expiry_date": format_user_date(product.expiry_date),
            }
            rows.append(
                {key: special.get(key, getattr(product, key, None)) for key in COLUMNS}
            )
        return rows

    @staticmethod
    def _cells(row: Dict[str, object]) -> List[object]:
        numbers = ("quantity", "min_quantity")
        return [
            format_quantity(float(row[key])) if key in numbers else row[key]
            for key in COLUMNS
        ]

    @staticmethod
    def _write_excel(path: Path, rows: List[Dict[str, object]]) -> None:
        from openpyxl import Workbook

        book = Workbook()
        sheet = book.active
        sheet.title = "Аптечка"
        sheet.append(list(COLUMNS.values()))
        for row in rows:
            sheet.append([row[key] for key in COLUMNS])
        for column, title in zip("ABCDEFGHI", COLUMNS.values()):
            sheet.column_dimensions[column].width = max(14, len(title) + 4)
        book.save(path)

    # --- шаблон для импорта ---

    def write_template(self, path: Union[str, Path], fmt: str) -> None:
        """Записывает шаблон для заполнения и последующего импорта.

        Excel: готовый файл программы с двумя таблицами («Моя аптечка» и
        «Список покупок»), списками категорий и единиц. CSV: заголовки столбцов
        и две строки-примера (только аптечка).

        Args:
            path: Куда сохранить.
            fmt: ``csv`` или ``xlsx``.

        Raises:
            ValidationError: Если формат не поддерживается.
        """
        if fmt not in ("csv", "xlsx"):
            raise ValidationError("Шаблон есть для файлов CSV и Excel")
        free_target(path)
        if fmt == "xlsx":
            shutil.copyfile(TEMPLATE_XLSX, path)
            return
        categories = [c.name for c in self._products.list_categories()]
        examples = [
            ["Ибупрофен", categories[0], 2, "упак.", "31.12.2027", "Жаропонижающее"]
            + ["Шкаф", 1, "Пример: замените или удалите строку"],
            ["Бинт", categories[1 % len(categories)], 5, "шт.", "", ""]
            + ["Аптечка", 2, ""],
        ]
        with Path(path).open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.writer(handle, delimiter=CSV_DELIMITER)
            writer.writerow(COLUMNS.values())
            writer.writerows(examples)

    # --- импорт ---

    def import_products(self, user_id: int, path: Union[str, Path]) -> ImportResult:
        """Добавляет товары (и позиции списка покупок) из файла CSV или Excel.

        Каждая строка проверяется так же, как форма добавления. Неверные строки
        пропускаются, остальные добавляются. В файле шаблона правая таблица
        («Товар», «Количество») попадает в список покупок.

        Args:
            user_id: Владелец товаров.
            path: Файл ``.csv`` или ``.xlsx``.

        Raises:
            ValidationError: Если формат файла не поддерживается или файл
                не удалось прочитать.
        """
        path = Path(path)
        suffix = path.suffix.lower().lstrip(".")
        if suffix not in ("csv", "xlsx"):
            raise ValidationError("Импорт возможен из файлов CSV и Excel")
        try:
            if suffix == "xlsx":
                records, wanted = self._read_xlsx(path)
            else:
                records, wanted = self._read_csv(path), []
        except (OSError, ValueError) as error:
            raise ValidationError(f"Не удалось прочитать файл: {error}")
        categories = {c.name.casefold(): c.id for c in self._products.list_categories()}
        result = ImportResult()
        for number, record in enumerate(records, start=1):
            try:
                self._products.add_product(user_id, self._form(record, categories))
                result.added += 1
            except ValidationError as error:
                row = record.get("_row", number)
                result.issues.append((KIT_SHEET, row, error.message))
        for number, record in enumerate(wanted, start=1):
            try:
                self._add_wanted(user_id, record)
                result.shopping_added += 1
            except ValidationError as error:
                row = record.get("_row", number)
                result.issues.append((SHOPPING_SHEET, row, error.message))
        return result

    def _add_wanted(self, user_id: int, record: Dict[str, object]) -> None:
        """Добавляет в список покупок строку правой таблицы шаблона."""
        if self._shopping is None:
            return
        quantity = _cell_text(record.get("quantity")) or "1"
        name = _cell_text(record.get("name"))
        self._shopping.add_manual(user_id, name, quantity, source=ShoppingSource.IMPORT)

    @staticmethod
    def _read_xlsx(
        path: Path,
    ) -> Tuple[List[Dict[str, object]], List[Dict[str, object]]]:
        """Читает файл Excel: товары и позиции списка покупок.

        В шаблоне товары лежат на листе «Моя аптечка», покупки на листе
        «Список покупок». В остальных файлах товары берутся с первого листа.
        """
        from openpyxl import load_workbook

        book = load_workbook(path, read_only=True, data_only=True)
        try:
            sheets = {
                sheet.title: list(sheet.iter_rows(values_only=True))
                for sheet in book.worksheets
            }
        finally:
            book.close()  # иначе Windows держит файл открытым
        first = next(iter(sheets.values()), [])
        kit = _table_records(sheets.get(KIT_SHEET, first))
        wanted = _table_records(sheets.get(SHOPPING_SHEET, []))
        return kit, wanted

    @staticmethod
    def _read_csv(path: Path) -> List[Dict[str, object]]:
        text = path.read_text(encoding="utf-8-sig")
        first = text.splitlines()[0] if text.strip() else ""
        delimiter = ";" if first.count(";") >= first.count(",") else ","
        reader = csv.DictReader(text.splitlines(), delimiter=delimiter)
        records = []
        for raw in reader:
            record = {}
            for title, value in raw.items():
                title = (title or "").strip()
                key = _BY_TITLE.get(_norm(title)) or (
                    title if title in COLUMNS else None
                )
                if key is not None:
                    record[key] = value
            records.append(record)
        return records

    @staticmethod
    def _form(record: Dict[str, object], categories: Dict[str, int]) -> ProductForm:
        def text(key: str) -> str:
            return _cell_text(record.get(key))

        category_id = categories.get(text("category").casefold())
        if category_id is None:
            raise ValidationError(f"неизвестная категория «{text('category')}»")
        texts = {key: text(key) for key in COLUMNS if key != "category"}
        return ProductForm(category_id=category_id, **texts)

    # --- сброс ---

    def reset_all(self, user_id: int) -> None:
        """Удаляет аптечку, список покупок, уведомления и историю пользователя.

        Сам аккаунт и его настройки остаются.
        """
        with self._db.transaction() as conn:
            for table in ("shopping_list", "notifications", "history", "products"):
                conn.execute(f"DELETE FROM {table} WHERE user_id = ?", (user_id,))
