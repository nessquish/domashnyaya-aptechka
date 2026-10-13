"""Собирает шаблон Excel для импорта: ``pharmacy/assets/template.xlsx``.

Запуск из папки проекта::

    .venv\\Scripts\\python.exe tools\\build_template.py

Три листа: «Правила» (как заполнять, категории и единицы), «Моя аптечка» и
«Список покупок». На листах с таблицами строка 1 — название, строка 2 —
заголовки: по ним импорт находит столбцы. Строки 3-6 оформлены заранее, а ниже
рамки и цвет продолжаются сами (условное форматирование).
"""

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from openpyxl import Workbook  # noqa: E402
from openpyxl.formatting.rule import FormulaRule  # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402

from pharmacy.db.schema import DEFAULT_CATEGORIES  # noqa: E402
from pharmacy.utils.units import UNITS  # noqa: E402

OUT = ROOT / "pharmacy" / "assets" / "template.xlsx"
LAST_ROW = 5000
PRE_ROWS = 6  # строки с готовым оформлением (3-6)

KIT = ("Название", "Категория", "Количество", "Ед.изм.", "Срок годности")
KIT += ("Место хранения",)
SHOPPING = ("Товар", "Количество")

EVEN_FILL = "FFFFFF"
# Цвета таблицы: (название, заголовки, нечётные строки).
PURPLE = ("B3A2C7", "CCC1DA", "E6E0EC")
BLUE = ("93CDDD", "B7DEE8", "DBEEF4")

RULES = (
    ("ПРАВИЛА ЗАПОЛНЕНИЯ", True),
    ("Лист «Моя аптечка»", True),
    ("• Обязательно: Название, Категория, Количество.", False),
    ("• Категория и Ед.изм.: выберите из списка в ячейке.", False),
    ("• Количество: число, дробное через запятую или точку.", False),
    ("• Срок годности: ДД.ММ.ГГГГ (например, 31.12.2027) или пусто.", False),
    ("• Дата добавления проставляется программой сама.", False),
    ("Лист «Список покупок»", True),
    ("• Обязательно: Товар. Количество по умолчанию 1.", False),
    ("Заполняйте строки подряд, без пропусков: рамки и цвет продолжаются.", False),
    ("Строки с ошибкой программа пропустит и покажет, какие именно.", False),
)
# Столбец -> (заголовок подсказки, текст подсказки) для каждой таблицы.
KIT_PROMPTS = {
    "A": ("Название", "Обязательно"),
    "C": ("Количество", "Обязательно. Число, например 2 или 1,5"),
    "E": ("Срок годности", "ДД.ММ.ГГГГ, например 31.12.2027, или пусто"),
}
SHOPPING_PROMPTS = {
    "A": ("Товар", "Обязательно"),
    "B": ("Количество", "Число. Если пусто, будет 1"),
}

side = Side(style="thin", color="FF000000")
BOX = Border(left=side, right=side, top=side, bottom=side)


def solid(color: str) -> PatternFill:
    return PatternFill(start_color=color, end_color=color, fill_type="solid")


def letter(column: int) -> str:
    return chr(ord("A") + column - 1)


def add_table(sheet, first: int, title: str, headers, colors=PURPLE) -> None:
    title_fill, head_fill, odd_fill = colors
    last = first + len(headers) - 1
    sheet.merge_cells(start_row=1, start_column=first, end_row=1, end_column=last)
    for column in range(first, last + 1):
        cell = sheet.cell(1, column)
        cell.fill, cell.border = solid(title_fill), BOX
    top = sheet.cell(1, first, title)
    top.font = Font(bold=True)
    top.alignment = Alignment(horizontal="center", vertical="center")
    for offset, text in enumerate(headers):
        cell = sheet.cell(2, first + offset, text)
        cell.font = Font(bold=True)
        cell.fill, cell.border = solid(head_fill), BOX
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in range(3, PRE_ROWS + 1):
        color = odd_fill if row % 2 else EVEN_FILL
        for column in range(first, last + 1):
            cell = sheet.cell(row, column)
            cell.fill, cell.border = solid(color), BOX
    area = f"{letter(first)}{PRE_ROWS + 1}:{letter(last)}{LAST_ROW}"
    has_data = f"COUNTA(${letter(first)}{PRE_ROWS + 1}:${letter(last)}{PRE_ROWS + 1})>0"
    for color, parity in ((odd_fill, 1), (EVEN_FILL, 0)):
        rule = FormulaRule(
            formula=[f"AND({has_data},MOD(ROW(),2)={parity})"],
            fill=solid(color),
            border=BOX,
        )
        sheet.conditional_formatting.add(area, rule)


def add_check(sheet, column: str, **options) -> None:
    check = DataValidation(allow_blank=True, **options)
    check.showInputMessage = True
    sheet.add_data_validation(check)
    check.add(f"{column}3:{column}{LAST_ROW}")


def style_lists(sheet) -> None:
    for check in sheet.data_validations.dataValidation:
        check.promptTitle, check.prompt = "Выберите из списка", "Стрелка справа"
        check.errorTitle, check.error = "Неверное значение", "Выберите из списка"
        check.showErrorMessage = True


def add_prompts(sheet, prompts) -> None:
    for column, (title, text) in prompts.items():
        add_check(sheet, column, type=None)
        check = sheet.data_validations.dataValidation[-1]
        check.promptTitle, check.prompt = title, text


def build_rules(sheet) -> None:
    sheet.column_dimensions["A"].width = 72
    sheet.column_dimensions["B"].width = 4
    sheet.column_dimensions["C"].width = 24
    sheet.column_dimensions["D"].width = 14
    for row, (text, bold) in enumerate(RULES, start=1):
        cell = sheet.cell(row, 1, text)
        cell.font = Font(bold=bold)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    sheet["C1"], sheet["D1"] = "Категории", "Единицы"
    sheet["C1"].font = sheet["D1"].font = Font(bold=True)
    for row, name in enumerate(DEFAULT_CATEGORIES, start=2):
        sheet.cell(row, 3, name)
    for row, unit in enumerate(UNITS, start=2):
        sheet.cell(row, 4, unit)


def build() -> None:
    book = Workbook()
    rules = book.active
    rules.title = "Правила"
    build_rules(rules)

    kit = book.create_sheet("Моя аптечка")
    add_table(kit, 1, "МОЯ АПТЕЧКА", KIT)
    for number, width in enumerate((24, 22, 14, 12, 16, 22), start=1):
        kit.column_dimensions[letter(number)].width = width
    # Срок годности вводится текстом ДД.ММ.ГГГГ: Excel не превратит его в число.
    for row in range(3, LAST_ROW + 1):
        kit[f"E{row}"].number_format = "@"
    # Списки: значения прямо в проверке, чтобы выпадающий список не был пустым.
    categories = ",".join(DEFAULT_CATEGORIES)
    add_check(kit, "B", type="list", formula1=f'"{categories}"')
    add_check(kit, "D", type="list", formula1='"' + ",".join(UNITS) + '"')
    style_lists(kit)
    add_prompts(kit, KIT_PROMPTS)

    shopping = book.create_sheet("Список покупок")
    add_table(shopping, 1, "СПИСОК ПОКУПОК", SHOPPING, BLUE)
    shopping.column_dimensions["A"].width = 30
    shopping.column_dimensions["B"].width = 14
    add_prompts(shopping, SHOPPING_PROMPTS)

    for sheet in (kit, shopping):
        sheet.freeze_panes = "A3"
    book.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
